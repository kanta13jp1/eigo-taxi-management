import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

OUTPUT_DIR = r"c:\Users\kanta\GitHub\eigo-taxi-management\docs\GoogleDrive_Upload\成果物一覧"
os.makedirs(OUTPUT_DIR, exist_ok=True)
FILE_PATH = os.path.join(OUTPUT_DIR, "Mauro_Brenda_Taxi_Link_成果物一覧.xlsx")

wb = openpyxl.Workbook()

# ==========================================
# 1. 成果物サマリー シート
# ==========================================
ws_summary = wb.active
ws_summary.title = "プロジェクト概要"
ws_summary.views.sheetView[0].showGridLines = True

font_title = Font(name="Meiryo", size=16, bold=True, color="059669")
font_section = Font(name="Meiryo", size=11, bold=True, color="0f172a")
font_header = Font(name="Meiryo", size=10, bold=True, color="FFFFFF")
font_body = Font(name="Meiryo", size=10, color="1e293b")
font_bold = Font(name="Meiryo", size=10, bold=True, color="1e293b")

fill_header = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
fill_sub = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

ws_summary["B2"] = "Mauro & Brenda Taxi Link 送迎チケット管理システム"
ws_summary["B2"].font = font_title
ws_summary["B3"] = "プロジェクト成果物一覧（Deliverables Inventory）"
ws_summary["B3"].font = font_section

proj_info = [
    ["プロジェクト名", "Mauro & Brenda Taxi Link (英語教室 送迎チケット管理)"],
    ["公開WebサイトURL", "https://kanta13jp1.github.io/eigo-taxi-management/"],
    ["ソースコードリポジトリ", "https://github.com/kanta13jp1/eigo-taxi-management"],
    ["主要関係者", "管理者: Mauro先生 & Brenda先生 / コーディネーター: 小林雅水様"],
    ["納品日 / バージョン", "2026年10月08日 / v1.0.0 (プロダクションリリース)"],
    ["開発技術スタック", "React 19, TypeScript, Tailwind CSS, Vite, GitHub Pages, Google Apps Script"],
    ["主要機能", "管理者・保護者の2軸UI、片道1枚消費ルール、フレックス予約、チケット残少連絡アラート、GAS自動連携"]
]

for row_idx, (k, v) in enumerate(proj_info, start=5):
    ws_summary.cell(row=row_idx, column=2, value=k).font = font_bold
    ws_summary.cell(row=row_idx, column=2).fill = fill_sub
    ws_summary.cell(row=row_idx, column=2).border = thin_border
    
    ws_summary.cell(row=row_idx, column=3, value=v).font = font_body
    ws_summary.cell(row=row_idx, column=3).border = thin_border

ws_summary.column_dimensions["B"].width = 22
ws_summary.column_dimensions["C"].width = 75

# ==========================================
# 2. 成果物一覧 シート
# ==========================================
ws_items = wb.create_sheet(title="成果物一覧")
ws_items.views.sheetView[0].showGridLines = True

headers = [
    "No", "成果物分類", "成果物名称", "形式 / 所在", 
    "概要・詳細説明", "主な利用者 / 用途", "ステータス"
]

for col_idx, h in enumerate(headers, start=1):
    cell = ws_items.cell(row=2, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border
ws_items.row_dimensions[2].height = 28

items_data = [
    # A. システム・アプリケーション本体
    [1, "システム本体", "送迎チケット管理Webアプリケーション", "Web (GitHub Pages)", "https://kanta13jp1.github.io/eigo-taxi-management/ にて稼働中。管理者用・送迎利用者用の2軸画面切替、完全予約制対応、チケット消化処理", "Mauro・Brenda先生、保護者、生徒", "完了（稼働中）"],
    [2, "システム本体", "フロントエンド ソースコード一式", "Git (GitHub)", "https://github.com/kanta13jp1/eigo-taxi-management にて公開管理。React 19, TypeScript, Tailwind CSS による型安全な堅牢実装", "開発者・運用保守", "完了（公開済）"],
    [3, "システム本体", "自動デプロイ CI/CD パイプライン", "GitHub Actions", "mainブランチへのプッシュで自動ビルド・自動本番公開を行うワークフロー定義 (.github/workflows/deploy.yml)", "開発者・運用保守", "完了（稼働中）"],
    
    # B. バックエンド・外部連携
    [4, "バックエンド・連携", "Google Apps Script (GAS) ソースコード", "JavaScript / GAS", "Googleスプレッドシート（3シート自動生成）およびGoogleカレンダーへの送迎予定自動作成・削除を行うバックエンドスクリプト", "システム運用者、先生", "完了（設定画面に同梱）"],
    [5, "バックエンド・連携", "Googleスプレッドシート連携仕様", "GSS定義", "「生徒台帳」「送迎履歴」「チケット購入履歴」の3テーブルスキーマ定義および自動データ同期インターフェース", "Mauro・Brenda先生、小林様", "完了"],
    [6, "バックエンド・連携", "Googleカレンダー連動機能", "Calendar API", "保護者または先生が予約確定時に「🚗【送迎】生徒名」の予定を自動作成、キャンセル時に自動削除する連動ロジック", "Mauro・Brenda先生", "完了"],

    # C. 設計・仕様書
    [7, "設計・仕様書", "画面デザイン仕様書 (PDF版)", "PDF (A4横)", "全10画面×PC版・スマホ版の【実機スクリーンショット計20枚】を収録したビジュアルUI/UXデザイン仕様書", "全関係者（先生・小林様・開発者）", "完了（PDF同梱）"],
    [8, "設計・仕様書", "画面デザイン仕様書 (Markdown版)", "Markdown (.md)", "デザインコンセプト、カラーパレット、レスポンシブ設計要件、ヒアリング反映事項を整理したドキュメント原本", "開発者・運用保守", "完了"],
    [9, "設計・仕様書", "要件定義 & ヒアリング反映仕様書", "ドキュメント", "「片道1枚消費」「完全予約制フレックス」「残数2回以下でMauro・Brendaに連絡してね！表示」「LINE/定期/決済の除外」の反映一覧", "小林様、先生方", "完了"],

    # D. テスト・品質検証
    [10, "テスト・品質検証", "受入シナリオテスト仕様書 (Excel版)", "Excel (.xlsx)", "保護者ポータル・先生ダッシュボード・Google連携・レスポンシブの全28項目テストマトリクス（合否チェックリスト付）", "テスター、小林様、開発者", "完了（Pass）"],
    [11, "テスト・品質検証", "受入シナリオテスト仕様書 (PDF版)", "PDF (A4横)", "全28項目の操作手順、前提条件、期待される結果、判定結果を美麗にスタイリングした配布・印刷用PDF", "全関係者", "完了（Pass）"],
    [12, "テスト・品質検証", "自動スクリーンショット & E2E検証スクリプト", "Node.js (Puppeteer)", "PC・スマホ両解像度での全画面自動キャプチャーおよびデータ同期検証を行う自動化スクリプト", "開発者", "完了（検証済）"],

    # E. 運用・ガイド
    [13, "運用マニュアル", "Google連携導入4ステップガイド", "Web内蔵ガイド", "Googleスプレッドシート新規作成からGASデプロイ、WebアプリURL貼り付けまでの手順をWeb画面内にわかりやすく内蔵", "Mauro・Brenda先生、小林様", "完了（Web画面内）"],
    [14, "運用マニュアル", "成果物一覧ファイル一式", "Excel / PDF / Markdown", "納品・共有対象となるすべての成果物とその所在・ステータスをまとめたインベントリ（本ファイル群）", "全関係者", "完了"]
]

for row_idx, row_data in enumerate(items_data, start=3):
    is_even = (row_idx % 2 == 0)
    row_fill = PatternFill(start_color="F8FAFC" if is_even else "FFFFFF", end_color="F8FAFC" if is_even else "FFFFFF", fill_type="solid")
    
    for col_idx, val in enumerate(row_data, start=1):
        cell = ws_items.cell(row=row_idx, column=col_idx, value=val)
        cell.font = font_body
        cell.fill = row_fill
        cell.border = thin_border
        
        if col_idx in [1, 2, 4, 7]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            
        if col_idx == 7 and "完了" in str(val):
            cell.font = Font(name="Meiryo", size=10, bold=True, color="059669")
            cell.fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")

    ws_items.row_dimensions[row_idx].height = 26

col_widths = {
    "A": 6,   # No
    "B": 18,  # 分類
    "C": 30,  # 名称
    "D": 18,  # 形式/所在
    "E": 48,  # 概要説明
    "F": 26,  # 主な利用者
    "G": 16   # ステータス
}
for col_letter, width in col_widths.items():
    ws_items.column_dimensions[col_letter].width = width

wb.save(FILE_PATH)
print(f"Successfully generated: {FILE_PATH}")
